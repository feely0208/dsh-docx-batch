#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
dsh-docx-batch 的真实处理逻辑（python-docx）

由 bin/cli.mjs 通过 spawnSync 调用：
    python3 lib/docx_tool.py <action> '<json payload>'
约定：
  - stdout 只输出一行 JSON（最后一行），人类可读文案由 Node 侧负责
  - stderr 输出异常栈（调试用）
  - 退出码：0 成功 / 1 运行失败 / 2 参数或输入问题 / 3 缺依赖
  - 只读输入文件，所有产物写进 --out 目录（默认 <输入目录>/_输出）
"""

import csv
import json
import os
import re
import sys
import tempfile
import traceback
from pathlib import Path

try:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
except ImportError as exc:  # 缺依赖必须明确提示安装命令，不能静默
    print(json.dumps({
        "ok": False,
        "code": 3,
        "message": "缺少 Python 依赖：%s" % getattr(exc, "name", exc),
        "hint": "请先安装：\n  python3 -m pip install python-docx\n"
                "如果 python3 也没有：\n  macOS: brew install python3\n  Ubuntu: sudo apt install python3",
    }, ensure_ascii=False))
    sys.exit(3)

DOCX_EXT = ".docx"
SKIP_DIRS = {"_输出", "node_modules", ".git", "__pycache__", ".venv", "venv"}
DEMO_HEADER = "机密 · 深鲸科技"
DEMO_FOOTER = "第 1 页"


# ── 错误 ──────────────────────────────────────────────────────────

class ToolError(Exception):
    def __init__(self, message, hint="", code=2):
        super().__init__(message)
        self.message = message
        self.hint = hint
        self.code = code


def emit(obj):
    print(json.dumps(obj, ensure_ascii=False))


# ── 路径与文件收集 ────────────────────────────────────────────────

def natural_key(text):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", text)]


def resolve_out(payload, input_path):
    """输出目录：--out 优先，否则 <输入目录>/_输出。"""
    out = payload.get("out")
    if out:
        return Path(out).expanduser().resolve()
    base = Path(input_path).expanduser().resolve()
    if not base.is_dir():
        base = base.parent
    return base / "_输出"


def _out_skip_set(out_dir):
    try:
        return {str(out_dir.resolve())}
    except OSError:
        return set()


def collect_docx(input_path, out_dir=None):
    """接受单个 .docx 或目录（递归），按文件名自然排序。"""
    if not input_path:
        raise ToolError("缺少 --input", "例：replace --input ./合同目录 --from 旧 --to 新")
    root = Path(input_path).expanduser().resolve()
    if not root.exists():
        raise ToolError(
            "输入不存在：%s" % input_path,
            "请检查路径是否正确；传目录会递归处理其中的 .docx 文件。",
        )
    skip_out = _out_skip_set(out_dir) if out_dir else set()
    if root.is_file():
        if root.suffix.lower() != DOCX_EXT:
            raise ToolError(
                "输入不是 .docx 文件：%s" % root.name,
                "本插件只处理 Word 2007+ 的 .docx；.doc 请先用 Word/WPS 另存为 .docx。",
            )
        return [root]
    files = []
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = sorted(
            [d for d in dirnames if d not in SKIP_DIRS and not d.startswith("~$")],
            key=natural_key,
        )
        if str(Path(dirpath).resolve()) in skip_out:
            dirnames[:] = []
            continue
        for name in filenames:
            if name.startswith("~$"):
                continue  # Word 打开时的临时文件
            p = Path(dirpath) / name
            if p.suffix.lower() == DOCX_EXT:
                files.append(p)
    files.sort(key=lambda p: natural_key(str(p.relative_to(root))))
    if not files:
        raise ToolError(
            "目录里没有 .docx 文件：%s" % root,
            "确认目录选对了；本插件不处理 .doc（请先另存为 .docx）。",
        )
    return files


def unique_path(directory, filename):
    """同名文件不覆盖，追加 _1/_2 序号。"""
    directory = Path(directory)
    target = directory / filename
    if not target.exists():
        return target
    stem, suffix = target.stem, target.suffix
    n = 1
    while (directory / ("%s_%d%s" % (stem, n, suffix))).exists():
        n += 1
    return directory / ("%s_%d%s" % (stem, n, suffix))


# ── docx 遍历 ─────────────────────────────────────────────────────

def iter_block_items(parent):
    """按文档顺序遍历段落和表格。"""
    body = parent.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def iter_cell_paragraphs(cell):
    for p in cell.paragraphs:
        yield p
    for table in cell.tables:
        for row in table.rows:
            for c in row.cells:
                yield from iter_cell_paragraphs(c)


def iter_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            yield from iter_cell_paragraphs(cell)


def all_editable_paragraphs(doc):
    """正文 + 表格 + 页眉页脚，去重（链接到前一节的页眉是同一个 part）。"""
    out, seen = [], set()

    def add(p):
        key = id(p._p)
        if key not in seen:
            seen.add(key)
            out.append(p)

    for p in doc.paragraphs:
        add(p)
    for t in doc.tables:
        for p in iter_table_paragraphs(t):
            add(p)
    for section in doc.sections:
        for part in (section.header, section.footer):
            try:
                for p in part.paragraphs:
                    add(p)
                for t in part.tables:
                    for p in iter_table_paragraphs(t):
                        add(p)
            except Exception:
                # 极少数残缺文档取不到页眉/页脚定义：跳过，不影响正文
                continue
    return out


# ── 文本替换 ──────────────────────────────────────────────────────

def replace_in_paragraph(paragraph, pairs):
    """先做 run 内替换（保留格式），跨 run 的再做段落级重建。返回替换次数。"""
    count = 0
    if not paragraph.runs:
        return 0
    for old, new in pairs:
        if not old:
            continue
        for run in paragraph.runs:
            if old in run.text:
                count += run.text.count(old)
                run.text = run.text.replace(old, new)
        if old in paragraph.text:
            runs = paragraph.runs
            full = "".join(r.text for r in runs)
            if old in full:
                count += full.count(old)
                runs[0].text = full.replace(old, new)
                for r in runs[1:]:
                    r.text = ""
    return count


# ── 各动作 ────────────────────────────────────────────────────────

def action_replace(payload):
    files = collect_docx(payload.get("input"))
    out = resolve_out(payload, payload.get("input"))
    olds = payload.get("from") or []
    news = payload.get("to") or []
    if not olds:
        raise ToolError("缺少 --from（要替换的旧文本）",
                        "例：replace --input ./合同 --from 旧公司名 --to 新公司名")
    if len(olds) != len(news):
        raise ToolError(
            "--from 与 --to 数量不一致（%d vs %d）" % (len(olds), len(news)),
            "每一组 --from 都要配一个 --to，按顺序一一对应。",
        )
    pairs = list(zip(olds, news))
    out.mkdir(parents=True, exist_ok=True)
    results, total = [], 0
    for src in files:
        doc = Document(str(src))
        n = 0
        for p in all_editable_paragraphs(doc):
            n += replace_in_paragraph(p, pairs)
        target = unique_path(out, src.name)
        doc.save(str(target))
        results.append({
            "input": str(src),
            "input_name": src.name,
            "output": str(target),
            "output_name": target.name,
            "replacements": n,
        })
        total += n
    return {
        "action": "replace",
        "input": str(Path(payload["input"]).expanduser().resolve()),
        "out": str(out),
        "files": results,
        "replacements": total,
        "summary": "处理 %d 个文件，共替换 %d 处" % (len(files), total),
        "artifacts": [r["output"] for r in results],
    }


def doc_body_text(doc):
    lines = []
    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            lines.append(block.text)
        else:
            lines.append("")
            lines.append("[表格]")
            for row in block.rows:
                lines.append("\t".join(c.text.replace("\n", " ") for c in row.cells))
    return "\n".join(lines).strip() + "\n"


def action_extract(payload):
    files = collect_docx(payload.get("input"))
    out = resolve_out(payload, payload.get("input"))
    out.mkdir(parents=True, exist_ok=True)
    results = []
    for src in files:
        text = doc_body_text(Document(str(src)))
        target = unique_path(out, src.stem + ".txt")
        target.write_text(text, encoding="utf-8")
        results.append({
            "input": str(src),
            "input_name": src.name,
            "output": str(target),
            "output_name": target.name,
            "chars": len(text),
        })
    return {
        "action": "extract",
        "input": str(Path(payload["input"]).expanduser().resolve()),
        "out": str(out),
        "files": results,
        "summary": "提取 %d 个文件的正文文本" % len(files),
        "artifacts": [r["output"] for r in results],
    }


def copy_run(dst_paragraph, src_run):
    run = dst_paragraph.add_run(src_run.text)
    try:
        run.bold = src_run.bold
        run.italic = src_run.italic
        run.underline = src_run.underline
        if src_run.font.size is not None:
            run.font.size = src_run.font.size
        if src_run.font.name:
            run.font.name = src_run.font.name
    except Exception:
        pass  # 字体属性个别无法复制时不影响文本合并


def copy_paragraph(dst_doc, src_paragraph):
    p = dst_doc.add_paragraph()
    try:
        style_name = src_paragraph.style.name if src_paragraph.style else None
        if style_name and style_name.lower() not in ("normal", "正文"):
            p.style = style_name
    except Exception:
        pass
    try:
        if src_paragraph.alignment is not None:
            p.alignment = src_paragraph.alignment
    except Exception:
        pass
    for run in src_paragraph.runs:
        copy_run(p, run)
    return p


def copy_table(dst_doc, src_table):
    rows = len(src_table.rows)
    cols = max((len(r.cells) for r in src_table.rows), default=0)
    if rows == 0 or cols == 0:
        return None
    table = dst_doc.add_table(rows=rows, cols=cols)
    for i, row in enumerate(src_table.rows):
        for j, cell in enumerate(row.cells):
            if j < cols:
                table.cell(i, j).text = cell.text
    try:
        table.style = src_table.style.name
    except Exception:
        pass
    return table


def action_merge(payload):
    files = collect_docx(payload.get("input"))
    out = resolve_out(payload, payload.get("input"))
    out.mkdir(parents=True, exist_ok=True)
    merged = Document()
    results = []
    for src in files:
        doc = Document(str(src))
        paragraphs_before = len(merged.paragraphs)
        for block in iter_block_items(doc):
            if isinstance(block, Paragraph):
                copy_paragraph(merged, block)
            else:
                copy_table(merged, block)
        results.append({
            "input": str(src),
            "input_name": src.name,
            "paragraphs": len(merged.paragraphs) - paragraphs_before,
            "source_paragraphs": len(doc.paragraphs),
        })
    target = unique_path(out, "merged.docx")
    merged.save(str(target))
    return {
        "action": "merge",
        "input": str(Path(payload["input"]).expanduser().resolve()),
        "out": str(out),
        "files": results,
        "output": str(target),
        "output_name": target.name,
        "paragraphs": len(merged.paragraphs),
        "summary": "按文件名顺序合并 %d 个文件，共 %d 段" % (len(files), len(merged.paragraphs)),
        "artifacts": [str(target)],
    }


def section_texts(doc):
    sections = []
    for idx, section in enumerate(doc.sections, 1):
        item = {"section": idx, "header": [], "footer": []}
        for label, part in (("header", section.header), ("footer", section.footer)):
            try:
                texts = [p.text.strip() for p in part.paragraphs if p.text.strip()]
            except Exception:
                texts = []
            item[label] = texts
        sections.append(item)
    return sections


def action_headers(payload):
    files = collect_docx(payload.get("input"))
    out = resolve_out(payload, payload.get("input"))
    out.mkdir(parents=True, exist_ok=True)
    results, lines = [], []
    for src in files:
        sections = section_texts(Document(str(src)))
        hit = sum(1 for s in sections if s["header"] or s["footer"])
        results.append({
            "input": str(src),
            "input_name": src.name,
            "sections": len(sections),
            "sections_with_text": hit,
            "detail": sections,
        })
        lines.append("文件：%s（%d 节）" % (src.name, len(sections)))
        for s in sections:
            h = " / ".join(s["header"]) if s["header"] else "（空）"
            f = " / ".join(s["footer"]) if s["footer"] else "（空）"
            lines.append("  第 %d 节  页眉：%s" % (s["section"], h))
            lines.append("          页脚：%s" % f)
        lines.append("")
    target = unique_path(out, "headers.txt")
    target.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")
    return {
        "action": "headers",
        "input": str(Path(payload["input"]).expanduser().resolve()),
        "out": str(out),
        "files": results,
        "output": str(target),
        "output_name": target.name,
        "summary": "列出 %d 个文件的页眉页脚" % len(files),
        "artifacts": [str(target)],
    }


def inline_runs(paragraph):
    parts = []
    for run in paragraph.runs:
        text = run.text
        if not text:
            continue
        if run.bold and run.italic:
            text = "***%s***" % text
        elif run.bold:
            text = "**%s**" % text
        elif run.italic:
            text = "*%s*" % text
        parts.append(text)
    text = "".join(parts).strip()
    return text or paragraph.text.strip()


def table_to_md(table):
    rows = []
    for row in table.rows:
        cells = [" ".join(c.text.split()) for c in row.cells]
        if any(cells):
            rows.append(cells)
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]

    def esc(s):
        return s.replace("|", "\\|")

    out = ["| " + " | ".join(esc(c) for c in rows[0]) + " |",
           "| " + " | ".join(["---"] * width) + " |"]
    for r in rows[1:]:
        out.append("| " + " | ".join(esc(c) for c in r) + " |")
    return "\n".join(out)


def paragraph_to_md(paragraph, state):
    style = (paragraph.style.name if paragraph.style else "") or ""
    style_low = style.lower()
    text = inline_runs(paragraph)
    m = re.match(r"^(heading|标题)\s*([1-6])$", style, re.I)
    if m:
        return "#" * int(m.group(2)) + " " + text
    if style_low.startswith("title"):
        return "# " + text
    if "list bullet" in style_low or ("列表" in style and "项目" in style):
        return "- " + text
    if "list number" in style_low:
        state["number"] += 1
        return "%d. %s" % (state["number"], text)
    if not text:
        return ""
    return text


def doc_to_md(doc):
    blocks = []
    state = {"number": 0}
    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            md = paragraph_to_md(block, state)
        else:
            md = table_to_md(block)
        if md.strip():
            blocks.append(md)
    return "\n\n".join(blocks).strip() + "\n"


def action_to_md(payload):
    files = collect_docx(payload.get("input"))
    out = resolve_out(payload, payload.get("input"))
    out.mkdir(parents=True, exist_ok=True)
    results = []
    for src in files:
        md = doc_to_md(Document(str(src)))
        target = unique_path(out, src.stem + ".md")
        target.write_text(md, encoding="utf-8")
        results.append({
            "input": str(src),
            "input_name": src.name,
            "output": str(target),
            "output_name": target.name,
            "chars": len(md),
        })
    return {
        "action": "to-md",
        "input": str(Path(payload["input"]).expanduser().resolve()),
        "out": str(out),
        "files": results,
        "summary": "转换 %d 个文件为 Markdown" % len(files),
        "artifacts": [r["output"] for r in results],
    }


# ── demo ─────────────────────────────────────────────────────────

def _make_demo_docx(path, heading, lines, table_rows=None, header=True):
    doc = Document()
    doc.add_heading(heading, level=1)
    for line in lines:
        doc.add_paragraph(line)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for i, row in enumerate(table_rows):
            for j, value in enumerate(row):
                table.cell(i, j).text = value
    if header:
        doc.sections[0].header.paragraphs[0].text = DEMO_HEADER
        doc.sections[0].footer.paragraphs[0].text = DEMO_FOOTER
    doc.save(str(path))


def action_demo(payload):
    root = Path(tempfile.mkdtemp(prefix="dsh-docx-batch-demo-"))
    inputs = root / "样例"
    inputs.mkdir(parents=True, exist_ok=True)
    _make_demo_docx(
        inputs / "合同A.docx",
        "第一章 总则",
        ["甲方：深鲸科技", "乙方：示例客户", "本合同由深鲸科技起草。"],
        [["条款", "内容"], ["保密", "深鲸科技对数据保密"]],
    )
    _make_demo_docx(
        inputs / "附件B.docx",
        "附件 B",
        ["附件内容第一段", "附件内容第二段"],
    )
    out = root / "输出"
    steps, artifacts = [], []

    def run(name, fn, extra):
        p = {"input": str(inputs), "out": str(out / name)}
        p.update(extra)
        result = fn(p)
        steps.append({
            "step": name,
            "action": result.get("action"),
            "summary": result.get("summary", ""),
            "out": result.get("out"),
        })
        artifacts.extend(result.get("artifacts", []))

    run("替换", action_replace, {"from": ["深鲸科技"], "to": ["深海鲸"]})
    replaced_dir = out / "替换"
    # 提取/转 Markdown 用替换后的产物，顺带演示链路
    p_extract = {"input": str(replaced_dir), "out": str(out / "文本")}
    result = action_extract(p_extract)
    steps.append({"step": "文本", "action": "extract", "summary": result["summary"], "out": result["out"]})
    artifacts.extend(result["artifacts"])
    p_md = {"input": str(replaced_dir), "out": str(out / "markdown")}
    result = action_to_md(p_md)
    steps.append({"step": "markdown", "action": "to-md", "summary": result["summary"], "out": result["out"]})
    artifacts.extend(result["artifacts"])
    run("合并", action_merge, {})
    run("页眉页脚", action_headers, {})

    return {
        "action": "demo",
        "demoRoot": str(root),
        "steps": steps,
        "summary": "演示完成：5 个动作全部跑通，产物 %d 个" % len(artifacts),
        "artifacts": artifacts,
    }


ACTIONS = {
    "replace": action_replace,
    "extract": action_extract,
    "merge": action_merge,
    "headers": action_headers,
    "to-md": action_to_md,
    "demo": action_demo,
}


def main(argv):
    if len(argv) < 3:
        emit({"ok": False, "code": 2,
              "message": "内部调用参数不足：需要 <action> 和 JSON payload",
              "hint": "请通过 bin/cli.mjs 调用本脚本。"})
        return 2
    action, raw = argv[1], argv[2]
    try:
        payload = json.loads(raw) if raw else {}
    except json.JSONDecodeError as exc:
        emit({"ok": False, "code": 2, "message": "参数 JSON 解析失败：%s" % exc, "hint": ""})
        return 2
    fn = ACTIONS.get(action)
    if not fn:
        emit({"ok": False, "code": 2, "message": "未知动作：%s" % action,
              "hint": "可用动作：%s" % " / ".join(k for k in ACTIONS if k != "demo")})
        return 2
    try:
        result = fn(payload)
        result.setdefault("ok", True)
        emit(result)
        return 0
    except ToolError as exc:
        emit({"ok": False, "code": exc.code, "message": exc.message, "hint": exc.hint})
        return exc.code
    except PermissionError as exc:
        emit({"ok": False, "code": 1,
              "message": "没有权限读写文件：%s" % exc,
              "hint": "检查目标文件是否被 Word 占用、或输出目录是否有写权限。"})
        return 1
    except Exception as exc:  # noqa: BLE001 —— 兜底，保证永远输出 JSON
        traceback.print_exc(file=sys.stderr)
        emit({"ok": False, "code": 1, "message": "处理失败：%s" % exc,
              "hint": "文件可能已损坏或被加密；可先用 Word 打开另存一次再试。"})
        return 1


if __name__ == "__main__":
    sys.exit(main(sys.argv))
